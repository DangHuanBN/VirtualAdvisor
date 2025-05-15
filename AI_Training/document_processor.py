# document_processor.py
import os
import PyPDF2
from docx import Document
from pptx import Presentation
import re

def process_document(file_path):
    """
    Xử lý các loại tài liệu khác nhau (PDF, DOCX, PPTX)
    và trả về nội dung dạng text
    """
    file_extension = os.path.splitext(file_path)[1].lower()
    
    if file_extension == '.pdf':
        return process_pdf(file_path)
    elif file_extension == '.docx':
        return process_docx(file_path)
    elif file_extension == '.pptx':
        return process_pptx(file_path)
    else:
        raise ValueError(f"Unsupported file format: {file_extension}")

def process_pdf(file_path):
    """Xử lý file PDF và trích xuất text"""
    text_content = []
    
    with open(file_path, 'rb') as file:
        pdf_reader = PyPDF2.PdfReader(file)
        num_pages = len(pdf_reader.pages)
        
        for page_num in range(num_pages):
            page = pdf_reader.pages[page_num]
            text = page.extract_text()
            if text:
                # Thêm thông tin trang để context rõ ràng hơn
                text_content.append(f"Page {page_num + 1}:\n{text}")
    
    return "\n\n".join(text_content)

def process_docx(file_path):
    """Xử lý file DOCX và trích xuất text"""
    doc = Document(file_path)
    
    # Lấy nội dung văn bản và đảm bảo giữ lại cấu trúc
    paragraphs = []
    for para in doc.paragraphs:
        if para.text.strip():
            paragraphs.append(para.text)
    
    # Xử lý các bảng
    tables = []
    for table in doc.tables:
        table_content = []
        for i, row in enumerate(table.rows):
            row_content = [cell.text for cell in row.cells]
            table_content.append(" | ".join(row_content))
        tables.append("\n".join(table_content))
    
    # Kết hợp nội dung
    content = "\n\n".join(paragraphs)
    if tables:
        content += "\n\nTABLES:\n" + "\n\n".join(tables)
    
    return content

def process_pptx(file_path):
    """Xử lý file PowerPoint và trích xuất text"""
    prs = Presentation(file_path)
    
    slides_content = []
    for i, slide in enumerate(prs.slides):
        slide_content = []
        slide_content.append(f"Slide {i+1}:")
        
        # Trích xuất tiêu đề
        if slide.shapes.title and slide.shapes.title.text:
            slide_content.append(f"Title: {slide.shapes.title.text}")
        
        # Trích xuất nội dung
        texts = []
        for shape in slide.shapes:
            if hasattr(shape, "text") and shape.text:
                text = shape.text.strip()
                if text and text not in texts:  # Tránh trùng lặp
                    texts.append(text)
        
        if texts:
            slide_content.append("Content:")
            slide_content.extend(texts)
        
        slides_content.append("\n".join(slide_content))
    
    return "\n\n".join(slides_content)

def chunk_text(text, chunk_size=1000, overlap=200):
    """
    Chia nhỏ văn bản thành các đoạn với kích thước cố định và có sự chồng lấp
    """
    chunks = []
    start = 0
    
    # Phân tách văn bản theo từng đoạn
    paragraphs = re.split(r'\n\s*\n', text)
    current_chunk = []
    current_size = 0
    
    for para in paragraphs:
        para_size = len(para)
        
        # Nếu đoạn hiện tại + đoạn mới vượt quá kích thước chunk
        if current_size + para_size > chunk_size and current_chunk:
            chunks.append("\n\n".join(current_chunk))
            # Giữ lại phần chồng lấp
            overlap_size = 0
            overlap_chunks = []
            for p in reversed(current_chunk):
                overlap_size += len(p)
                overlap_chunks.insert(0, p)
                if overlap_size >= overlap:
                    break
            
            current_chunk = overlap_chunks
            current_size = overlap_size
        
        current_chunk.append(para)
        current_size += para_size
    
    # Thêm chunk cuối cùng
    if current_chunk:
        chunks.append("\n\n".join(current_chunk))
    
    return chunks